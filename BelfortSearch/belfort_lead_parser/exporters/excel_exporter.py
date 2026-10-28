"""Экспорт в Excel (.xlsx) через openpyxl (ТЗ п.14)."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter

from exporters.csv_exporter import CSV_COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_excel(rows: Iterable, path: str | Path, stats: dict | None = None,
                 doubtful_rows: Iterable | None = None) -> Path:
    """Создаёт .xlsx с фильтрами, закреплённой шапкой и кликабельными ссылками."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"

    # Шапка.
    headers = [title for _, title in CSV_COLUMNS]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = "A2"

    # Данные.
    data_start = 2
    for offset, row in enumerate(rows):
        r = data_start + offset
        for col_idx, (field_name, _) in enumerate(CSV_COLUMNS, start=1):
            value = _get(row, field_name)
            cell = ws.cell(row=r, column=col_idx, value="" if value is None else value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if field_name in ("card_url", "website_url") and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
        if offset == 0:
            pass

    # Автофильтр.
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # Ширина колонок.
    widths = {
        "ID": 10, "Название компании": 30, "Ниша": 18, "Город": 18, "Адрес": 32,
        "Телефон": 18, "Ссылка на карточку": 30, "Источник": 16, "Рейтинг": 8,
        "Отзывы": 8, "Статус сайта": 22, "Балл": 7, "Расшифровка балла": 50,
        "Статус обработки": 16, "Дата первого обнаружения": 22,
        "Дата последней проверки": 22,
    }
    for col_idx, title in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(title, 14)

    # Лист статистики запуска.
    if stats is not None:
        stats_ws = wb.create_sheet("Статистика")
        stats_ws.append(["Показатель", "Значение"])
        for key, value in stats.items():
            stats_ws.append([key, value])
        for col in ("A", "B"):
            stats_ws.column_dimensions[col].width = 32

    # Лист сомнительных совпадений.
    if doubtful_rows is not None:
        doubtful_list = list(doubtful_rows)
        if doubtful_list:
            dws = wb.create_sheet("Сомнительные совпадения")
            dws.append(headers)
            for offset, row in enumerate(doubtful_list):
                r = 2 + offset
                for col_idx, (field_name, _) in enumerate(CSV_COLUMNS, start=1):
                    value = _get(row, field_name)
                    dws.cell(row=r, column=col_idx, value="" if value is None else value)
            dws.freeze_panes = "A2"
            dws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{dws.max_row}"

    wb.save(path)
    return path


def _get(row, field_name: str):
    if isinstance(row, dict):
        return row.get(field_name)
    return getattr(row, field_name, None)
